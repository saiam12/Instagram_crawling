from __future__ import annotations

import csv
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from android_collector.models import EvidencePaths, Metric, ObservedProfile, ObservedReel
from android_collector.store import (
    ROW_COLLECTION_FIELDS,
    USER_COLLECTION_FIELDS,
    CollectionStore,
    read_reel_urls_from_xlsx,
)


class FakeDriver:
    def capture_screenshot(self, path: Path) -> None:
        path.write_bytes(b"\x89PNG\r\n\x1a\n")


def observed_reel(evidence_paths: EvidencePaths) -> ObservedReel:
    return ObservedReel(
        source_mode="feed",
        source_query="",
        reel_url="",
        reel_fingerprint="fingerprint",
        collected_at="2026-08-30T00:00:00Z",
        username="odi.pigi",
        caption="caption",
        audio_name="audio",
        uploaded_at="2026-04-28",
        profile=ObservedProfile(
            biography="profile bio",
            profile_category="Digital creator",
            account_country="Argentina",
            post_count=7,
            following_count=321,
            follower_count=1_169,
        ),
        metrics={
            "likes_and_plays_count": Metric("Likes and plays", 32_357, "32,357"),
            "like_count": Metric("like", 4_699, "4,699"),
        },
        visible_metrics={"Likes and plays": "32,357"},
        evidence_paths=evidence_paths,
    )


class CollectionStoreTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.data_dir = Path(self.temporary.name)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def test_export_writes_matching_csv_json_and_xlsx(self) -> None:
        store = CollectionStore(self.data_dir)
        evidence = store.save_evidence(1, "<hierarchy/>", FakeDriver())
        store.append(observed_reel(evidence))

        store.export()

        with (self.data_dir / "reels.csv").open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            self.assertEqual(reader.fieldnames, list(ROW_COLLECTION_FIELDS))
            row = next(reader)
        payload = json.loads((self.data_dir / "reels.json").read_text(encoding="utf-8"))
        self.assertEqual(row["collection_number"], "1")
        self.assertEqual(row["title"], "caption")
        self.assertEqual(row["view_count"], "")
        self.assertEqual(row["like_count"], "4699")
        self.assertEqual(payload[0]["like_count"], 4_699)
        self.assertEqual(payload[0]["uploaded_at"], "2026-04-28")
        self.assertEqual(payload[0]["days_since_upload"], 124)
        self.assertIsNone(payload[0]["view_count"])
        self.assertEqual(list(payload[0]), list(ROW_COLLECTION_FIELDS))
        raw_payload = json.loads(
            (self.data_dir / ".collector" / "android_observations.json").read_text(encoding="utf-8")
        )
        self.assertEqual(raw_payload[0]["likes_and_plays_count"], 32_357)
        self.assertEqual(raw_payload[0]["likes_and_plays_count_raw"], "32,357")
        self.assertEqual(raw_payload[0]["evidence_xml_path"], evidence.xml_path)
        workbook = load_workbook(self.data_dir / "reels.xlsx", data_only=True)
        try:
            self.assertEqual(
                [cell.value for cell in next(workbook.active.iter_rows(max_row=1))],
                list(ROW_COLLECTION_FIELDS),
            )
            self.assertEqual(workbook.active.freeze_panes, "A2")
            self.assertEqual(workbook.active["A1"].fill.fgColor.rgb, "000F766E")
            self.assertEqual(workbook.active.auto_filter.ref, f"A1:Z2")
            self.assertEqual(workbook.active["L2"].value.date().isoformat(), "2026-04-28")
            self.assertEqual(workbook.active["L2"].number_format, "yyyy-mm-dd")
            self.assertEqual(workbook.active["N2"].value, 124)
        finally:
            workbook.close()

        users_payload = json.loads((self.data_dir / "users.json").read_text(encoding="utf-8"))
        self.assertEqual(list(users_payload[0]), list(USER_COLLECTION_FIELDS))
        self.assertEqual(users_payload[0]["username"], "odi.pigi")
        self.assertEqual(users_payload[0]["biography"], "profile bio")
        self.assertEqual(users_payload[0]["profile_category"], "Digital creator")
        self.assertEqual(users_payload[0]["post_count"], 7)
        self.assertEqual(users_payload[0]["following_count"], 321)
        self.assertEqual(users_payload[0]["follower_count"], 1_169)
        self.assertEqual(users_payload[0]["account_country"], "Argentina")
        users_workbook = load_workbook(self.data_dir / "users.xlsx", read_only=True, data_only=True)
        try:
            self.assertEqual(
                [cell.value for cell in next(users_workbook.active.iter_rows(max_row=1))],
                list(USER_COLLECTION_FIELDS),
            )
            self.assertEqual(users_workbook.active["H2"].value, 321)
            self.assertEqual(users_workbook.active["H2"].number_format, "#,##0")
        finally:
            users_workbook.close()
        self.assertTrue(Path(evidence.xml_path).exists())
        self.assertTrue(Path(evidence.png_path).exists())

    def test_xml_only_evidence_avoids_the_expensive_screenshot_round_trip(self) -> None:
        store = CollectionStore(self.data_dir)

        evidence = store.save_evidence(
            1,
            "<hierarchy><node text=\"Reels\" /></hierarchy>",
            FakeDriver(),
            capture_screenshot=False,
        )

        self.assertTrue(Path(evidence.xml_path).exists())
        self.assertEqual(evidence.png_path, "")
        self.assertFalse((self.data_dir / "evidence" / "000001.png").exists())

    def test_xlsx_url_reader_rejects_non_instagram_rows(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["url"])
        worksheet.append(["https://www.instagram.com/reel/ABC/"])
        worksheet.append(["https://example.com/"])
        path = self.data_dir / "urls.xlsx"
        workbook.save(path)

        self.assertEqual(read_reel_urls_from_xlsx(path), ["https://www.instagram.com/reel/ABC/"])

    def test_xlsx_url_reader_limits_visible_excel_rows_inclusively(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["name", "url"])
        worksheet.append(["first", "https://www.instagram.com/reel/ROW2/"])
        worksheet.append(["second", "https://www.instagram.com/reel/ROW3/"])
        worksheet.append(["third", "https://www.instagram.com/reel/ROW4/"])
        path = self.data_dir / "selected_rows.xlsx"
        workbook.save(path)
        workbook.close()

        self.assertEqual(
            read_reel_urls_from_xlsx(path, start_row=3, end_row=4),
            [
                "https://www.instagram.com/reel/ROW3/",
                "https://www.instagram.com/reel/ROW4/",
            ],
        )

    def test_user_history_keeps_profile_fields_and_calculates_follower_change(self) -> None:
        store = CollectionStore(self.data_dir)
        evidence = store.save_evidence(1, "<hierarchy/>", FakeDriver())
        first = observed_reel(evidence)
        second = replace(
            first,
            collected_at="2026-08-30T01:00:00Z",
            profile=replace(first.profile, follower_count=1_172),
        )
        store.append(first)
        store.append(second)

        store.export()

        users = json.loads((self.data_dir / "users.json").read_text(encoding="utf-8"))
        self.assertEqual(users[0]["biography"], "profile bio")
        self.assertEqual(users[0]["profile_category"], "Digital creator")
        self.assertEqual(users[0]["post_count"], 7)
        self.assertEqual(users[0]["following_count"], 321)
        self.assertEqual(users[0]["follower_count"], 1_169)
        self.assertEqual(users[0]["account_country"], "Argentina")
        self.assertIsNone(users[0]["follower_count_change"])
        self.assertEqual(users[1]["collection_number"], 2)
        self.assertEqual(users[1]["follower_count_change"], 3)

    def test_url_refresh_identity_ignores_shared_link_query_parameters(self) -> None:
        store = CollectionStore(self.data_dir)
        evidence = store.save_evidence(1, "<hierarchy/>", FakeDriver())
        first = replace(
            observed_reel(evidence),
            reel_url="https://www.instagram.com/reel/CODE123/?igsh=first",
            collected_at="2026-08-30T00:00:00Z",
        )
        second = replace(
            first,
            reel_url="https://www.instagram.com/reel/CODE123/?igsh=second",
            collected_at="2026-08-30T01:00:00Z",
            metrics={**first.metrics, "like_count": Metric("like", 4_700, "4700")},
        )
        store.append(first)
        store.append(second)
        store.export()

        reels = json.loads((self.data_dir / "reels.json").read_text(encoding="utf-8"))

        self.assertEqual(reels[0]["collection_number"], 1)
        self.assertEqual(reels[1]["collection_number"], 2)

    def test_public_ad_field_marks_explicit_ads_or_sponsored_hashtags(self) -> None:
        store = CollectionStore(self.data_dir)
        evidence = store.save_evidence(1, "<hierarchy/>", FakeDriver())
        sponsored = replace(observed_reel(evidence), caption="provided #협찬")
        explicit_ad = replace(
            observed_reel(evidence),
            reel_fingerprint="explicit-ad",
            username="advertiser",
            is_ad=True,
        )
        store.append(sponsored)
        store.append(explicit_ad)

        store.export()

        reels = json.loads((self.data_dir / "reels.json").read_text(encoding="utf-8"))
        self.assertEqual(reels[0]["ad"], "협찬")
        self.assertTrue(reels[1]["ad"])

    def test_export_migrates_legacy_android_public_file_without_losing_raw_fields(self) -> None:
        legacy = [
            {
                "collected_at": "2026-08-30T00:00:00Z",
                "source_mode": "hashtag",
                "source_query": "패션",
                "reel_url": "",
                "reel_fingerprint": "legacy-fingerprint",
                "username": "creator",
                "caption": "style #ootd",
                "audio_name": "audio",
                "like_count": 7,
                "likes_and_plays_count": 99,
                "likes_and_plays_count_raw": "99",
                "evidence_xml_path": "evidence/000001.xml",
                "evidence_png_path": "evidence/000001.png",
                "status": "collected",
            }
        ]
        (self.data_dir / "reels.json").write_text(json.dumps(legacy), encoding="utf-8")

        CollectionStore(self.data_dir).export()

        public = json.loads((self.data_dir / "reels.json").read_text(encoding="utf-8"))
        internal = json.loads(
            (self.data_dir / ".collector" / "android_observations.json").read_text(encoding="utf-8")
        )
        self.assertEqual(list(public[0]), list(ROW_COLLECTION_FIELDS))
        self.assertEqual(public[0]["like_count"], 7)
        self.assertEqual(public[0]["hashtags"], "#ootd")
        self.assertEqual(internal[0]["source_query"], "패션")
        self.assertEqual(internal[0]["likes_and_plays_count"], 99)

    def test_export_backfills_legacy_audio_from_visible_evidence(self) -> None:
        store = CollectionStore(self.data_dir)
        evidence_xml = """
        <hierarchy>
          <node text="odi.pigi" resource-id="com.instagram.android:id/clips_author_username" />
          <node content-desc="Snoh Aalegra · Nothing Burns Like The Cold" />
          <node resource-id="com.instagram.android:id/clips_caption_component" />
        </hierarchy>
        """
        evidence = store.save_evidence(1, evidence_xml, FakeDriver())
        store.append(replace(observed_reel(evidence), audio_name=""))

        store.export()

        reels = json.loads((self.data_dir / "reels.json").read_text(encoding="utf-8"))
        raw = json.loads((self.data_dir / ".collector" / "android_observations.json").read_text(encoding="utf-8"))
        self.assertEqual(reels[0]["audio_name"], "Snoh Aalegra · Nothing Burns Like The Cold")
        self.assertEqual(raw[0]["audio_name"], "Snoh Aalegra · Nothing Burns Like The Cold")


if __name__ == "__main__":
    unittest.main()
