"""Reel URL 입력 소스 처리."""

import re
from pathlib import Path


SHORTCODE_PATTERN = re.compile(r"^[A-Za-z0-9_-]+$")

def is_instagram_reel_url(value: object) -> bool:
    url = str(value or "").strip().lower()
    prefixes = (
        "https://www.instagram.com/reel/",
        "https://www.instagram.com/reels/",
        "https://instagram.com/reel/",
        "https://instagram.com/reels/",
        "https://instagram.com/p/",
    )
    return url.startswith(prefixes)


def normalize_reel_url(value: object) -> str:
    text = str(value or "").strip()
    if is_instagram_reel_url(text):
        return text
    if SHORTCODE_PATTERN.fullmatch(text):
        return f"https://www.instagram.com/reels/{text}/"
    return ""


def read_reel_urls_from_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise ValueError(f"XLSX 파일을 확인하세요: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    urls, seen = [], set()
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue
            names = [str(value or "").strip().lower() for value in headers]
            url_index = next((i for i, name in enumerate(names) if name in {"url", "reel_url"}), None)
            if url_index is None:
                continue
            for row in rows:
                value = row[url_index] if url_index < len(row) else None
                url = normalize_reel_url(value)
                if url and url not in seen:
                    seen.add(url)
                    urls.append(url)
    finally:
        workbook.close()
    return urls
