import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from reel_analyzer import GroupedVideoAnalyses, VideoAnalysis, _append_readable_output


def sample_analysis():
    return {
        "summary": "두 사람이 각자의 코디를 보여준다.",
        "hook": {"description": "두 사람이 등장한다.", "strength": "medium"},
        "body_structure": "착장을 보여준다.",
        "camera": {"main_composition": "전신", "angles": [], "movements": []},
        "editing": {"style": "하드 컷", "cut_speed": "medium", "transitions": []},
        "subtitle": {"exists": False, "position": None, "style": None},
        "subjects": {
            "people_count": 2,
            "main_subject": "인물 A와 인물 B",
            "people": [
                {"subject_id": "인물 A", "gender_presentation": "male"},
                {"subject_id": "인물 B", "gender_presentation": "female"},
            ],
        },
        "product": {"exists": True, "description": "재킷과 코트"},
        "scene_details": [{
            "scene_number": 1, "start_second": 0.0, "end_second": 2.0, "section": "hook",
            "visual_description": "두 사람이 나란히 서 있다.",
            "worn_outfits": [
                {"subject_id": "인물 A", "gender_presentation": "male", "clothing_items": ["검은 재킷"]},
                {"subject_id": "인물 B", "gender_presentation": "female", "clothing_items": ["흰 코트"]},
            ],
            "layout_type": "full_body_fashion_shot", "camera": "전신",
            "camera_movement": "없음", "on_screen_text": "20대 남녀 코디",
            "audio_or_dialogue": None, "transition_in": "none", "emotional_tone": "밝음",
            "purpose": "착장 소개",
        }],
        "audio_analysis": {
            "speech_present": False, "language": None, "transcript": [], "sound_events": [],
            "background_music": {"exists": False, "mood": None, "tempo": "none", "vocals": None},
        },
        "content_type": "lookbook",
        "thumbnail_analysis": {
            "source": "video_first_frame", "visual_description": "두 사람", "on_screen_text": None,
            "focal_point": "두 사람", "composition": "전신", "selling_point": None,
            "strengths": [], "weaknesses": [], "effectiveness": "medium", "evidence_confidence": 0.8,
        },
        "marketing_analysis": {"strengths": [], "weaknesses": [], "notable_elements": [], "selling_points": []},
        "recommended_audience": [
            {"age_group": "20s", "gender": "male", "evidence": "화면 문구: 20대 남녀 코디"},
            {"age_group": "20s", "gender": "female", "evidence": "화면 문구: 20대 남녀 코디"},
        ],
        "generation_prompts": {
            "video_prompt_en": "Two people show outfits.",
            "graphic_post_processing_needed": False, "post_processing_notes": None,
        },
    }


class AudienceSchemaTests(unittest.TestCase):
    def test_people_outfits_and_multiple_audiences_survive_json_validation(self):
        result = VideoAnalysis.model_validate_json(json.dumps(sample_analysis(), ensure_ascii=False)).model_dump()
        self.assertEqual(result["subjects"]["people"][1]["gender_presentation"], "female")
        self.assertEqual(result["scene_details"][0]["worn_outfits"][0]["subject_id"], "인물 A")
        self.assertEqual(len(result["recommended_audience"]), 2)
        grouped = GroupedVideoAnalyses.model_validate({
            "analyses": [{"input_index": 1, "analysis": sample_analysis()}]
        })
        self.assertEqual(grouped.analyses[0].analysis.recommended_audience[0].age_group, "20s")

    def test_flat_lay_has_no_wearer_and_can_use_broad_audience(self):
        payload = sample_analysis()
        payload["subjects"].update(people_count=0, people=[])
        payload["scene_details"][0]["worn_outfits"] = []
        payload["recommended_audience"] = [{
            "age_group": "20s_30s", "gender": "female", "evidence": "20~30대 여성 코디 문구",
        }]
        result = VideoAnalysis.model_validate(payload)
        self.assertEqual(result.scene_details[0].worn_outfits, [])
        self.assertEqual(result.recommended_audience[0].age_group, "20s_30s")

    def test_rejects_unlinked_or_inconsistent_outfit(self):
        for subject_id, gender in (("인물 C", "male"), ("인물 A", "female")):
            with self.subTest(subject_id=subject_id, gender=gender):
                payload = sample_analysis()
                payload["scene_details"][0]["worn_outfits"][0].update(
                    subject_id=subject_id, gender_presentation=gender
                )
                with self.assertRaises(ValidationError):
                    VideoAnalysis.model_validate(payload)

    def test_existing_xlsx_gets_audience_columns_without_losing_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "reel_analyses.json"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([
                "reel_id", "analyzed_at", "summary", "hook", "hook_strength", "body_structure",
                "content_type", "scene_count", "transcript", "background_music", "selling_points",
            ])
            sheet.append(["old-reel"])
            workbook.save(output.with_suffix(".xlsx"))
            workbook.close()
            with patch("reel_analyzer.OUTPUT_FILE", str(output)):
                _append_readable_output("new-reel", "2026-09-24", sample_analysis())
            workbook = load_workbook(output.with_suffix(".xlsx"), read_only=True)
            sheet = workbook.active
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["A2"].value, "old-reel")
            self.assertEqual(sheet["L1"].value, "subject_genders")
            self.assertIn("인물 A", sheet["M3"].value)
            self.assertIn("20대 남성", sheet["N3"].value)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
